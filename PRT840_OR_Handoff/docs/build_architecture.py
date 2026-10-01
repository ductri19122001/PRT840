from pathlib import Path
from html import escape
import os
import subprocess
import json

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
OUT = BASE
OUT.mkdir(parents=True, exist_ok=True)
SVG = OUT / 'PRT840_System_Architecture.svg'
PNG = OUT / 'PRT840_System_Architecture.png'
DOCX = OUT / 'PRT840_System_Architecture.docx'

parts = ['<svg xmlns="http://www.w3.org/2000/svg" width="1800" height="1160" viewBox="0 0 1800 1160">',
         '<title>PRT840 network detection and LLM assisted evidence interpretation architecture</title>',
         '<desc>Saved Decision Tree and Isolation Forest predictions are joined by dataset and flow identifier using a three-state OR rule. Positive alerts open a declared endpoint and date scope. Trigger alerts and wider processed flow context are kept separate. Network evidence, retrospective CTI and official ATTACK references enter a local Ollama request. Generated drafts require evidence review. The full saved VM3 Decision Tree test export is included; detector test populations differ.</desc>',
         '<defs><marker id="arrow" markerUnits="userSpaceOnUse" markerWidth="15" markerHeight="15" refX="14" refY="7.5" orient="auto"><path d="M0 0 L15 7.5 L0 15 Z" fill="#536879"/></marker><marker id="dasharrow" markerUnits="userSpaceOnUse" markerWidth="15" markerHeight="15" refX="14" refY="7.5" orient="auto"><path d="M0 0 L15 7.5 L0 15 Z" fill="#A15D16"/></marker></defs>',
         '<rect width="1800" height="1160" fill="white"/>']

def text(x,y,s,size=26,weight=400,colour='#19232B',anchor='start'):
    parts.append(f'<text x="{x}" y="{y}" font-family="DejaVu Sans, Arial, sans-serif" font-size="{size}" font-weight="{weight}" fill="{colour}" text-anchor="{anchor}">{escape(s)}</text>')

def line(coords, dashed=False):
    pts=' '.join(f'{x},{y}' for x,y in coords)
    style='stroke-dasharray="11 8"' if dashed else ''
    colour='#A15D16' if dashed else '#536879'
    marker='dasharrow' if dashed else 'arrow'
    parts.append(f'<polyline points="{pts}" fill="none" stroke="{colour}" stroke-width="3" stroke-linejoin="round" {style} marker-end="url(#{marker})"/>')

def node(x,y,w,h,title,lines,fill='#F3F6F8'):
    parts.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="9" fill="{fill}" stroke="#889AA7" stroke-width="1.7"/>')
    text(x+w/2,y+39,title,28,600,anchor='middle')
    first = 70 if len(lines)>2 else 76
    for i,s in enumerate(lines):
        text(x+w/2,y+first+29*i,s,23,anchor='middle')

def phase(y,number,title):
    text(55,y,f'{number}  {title}',29,600)

phase(34,'1','Saved behavioural detection evidence')
node(55,90,260,150,'Zeek logs',['VM2 / VM3 records','Connection records'])
node(370,90,300,150,'Data preparation',['Behavioural feature table','Separate flow metadata'])
node(730,65,320,110,'Decision Tree',['Saved test predictions'])
node(730,230,320,110,'Isolation Forest',['Saved anomaly decisions'])
node(1115,110,260,180,'OR rule',['Either flags: alert','Both normal: no alert','Otherwise: unknown'])
node(1450,120,295,150,'Positive alerts',['Flow identifiers','Model decisions retained'])
line([(315,165),(370,165)])
line([(670,165),(700,165),(700,120),(730,120)])
line([(670,165),(700,165),(700,285),(730,285)])
line([(1050,120),(1080,120),(1080,153),(1115,153)])
line([(1050,285),(1080,285),(1080,247),(1115,247)])
line([(1375,195),(1450,195)])
line([(510,240),(510,262),(350,262),(350,352),(25,352),(25,517),(55,517)])
text(55,333,'Context flow records',22)
text(380,302,'Saved model exports',22)
text(380,333,'No model refitting',22)

phase(410,'2','Evidence and LLM interpretation')
node(55,445,320,145,'Case preparation',['Endpoint and date scope','Positive alert required'])
node(470,445,330,145,'Network evidence',['Triggers and wider context','Counts, states and traffic'])
node(935,445,320,145,'CTI assessment',['Endpoint relationships','Sources and limitations'])
node(1430,445,320,145,'Public CTI sources',['MalwareBazaar','VirusTotal','Retrospective findings'])
line([(375,517),(470,517)])
line([(800,517),(935,517)])
line([(1430,517),(1255,517)])
line([(1598,270),(1598,424),(215,424),(215,445)])
text(850,400,'Positive alerts open the declared case scope',22)

node(55,650,320,145,'Official ATT&CK',['Pinned MITRE definitions','Reference collection'])
node(470,650,330,145,'Reference selection',['BM25 keyword retrieval','Candidate definitions'])
node(935,650,320,145,'Prompt assembly',['Network evidence and CTI','ATT&CK references','Instructions and schema'])
node(1430,650,320,145,'Local LLM',['Llama 3.1 8B via Ollama','Python calls the local API'])
line([(375,722),(470,722)])
line([(800,722),(935,722)])
line([(1095,590),(1095,650)])
line([(1255,722),(1430,722)])

phase(852,'3','Human review and reporting')
node(1430,895,320,135,'Human review',['Check facts and mappings','Review response guidance'])
node(935,895,320,135,'Assessment record',['Supported / unsupported /','uncertain claims'])
node(470,895,330,135,'Case study report',['Evidence and findings','Errors and limitations'])
line([(1590,795),(1590,895)])
text(1610,851,'Draft answer',23)
line([(1430,962),(1255,962)])
line([(935,962),(800,962)])
line([(60,925),(150,925)])
text(170,933,'Architecture data flow',22)
text(55,982,'VM3 DT coverage:',22)
text(55,1014,'27,953 saved test rows',22)
text(55,1088,'Retrospective replay: dataset labels are withheld; alert counts and endpoint context enter the prompt.',23)
text(55,1128,'Historical preselected-case experiments remain separate. Generated interpretations require evidence review.',23)
parts.append('</svg>')
SVG.write_text('\n'.join(parts),encoding='utf-8')

node_js = "const sharp=require(process.env.CODEX_PRIMARY_RUNTIME_NODE_MODULES+'/sharp'); sharp(process.argv[1], {density:180}).resize({width:3600}).png().toFile(process.argv[2]).catch(e=>{process.stderr.write(String(e));process.exit(1)});"
subprocess.run([os.environ['CODEX_PRIMARY_RUNTIME_NODE'],'-e',node_js,str(SVG),str(PNG)],check=True)

doc=Document()
sec=doc.sections[0]
sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width=Inches(11)
sec.page_height=Inches(8.5)
sec.top_margin=Inches(.48)
sec.bottom_margin=Inches(.48)
sec.left_margin=Inches(.65)
sec.right_margin=Inches(.65)
sec.header_distance=Inches(.20)
sec.footer_distance=Inches(.20)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    s=doc.styles[name]
    s.font.name='Calibri'
    s.font.color.rgb=RGBColor(0,0,0)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(23)
doc.styles['Title'].paragraph_format.space_after=Pt(5)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(9)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(6)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(5)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].font.bold=False
doc.styles['Caption'].paragraph_format.space_after=Pt(4)

header=sec.header.paragraphs[0]
header.text='PRT840    System architecture    1 October 2026'
header.runs[0].font.size=Pt(8)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run()
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)

doc.add_paragraph('Network evidence interpretation architecture','Title')
p=doc.add_paragraph('Behavioural detection, CTI context and LLM assisted forensic interpretation')
p.paragraph_format.space_after=Pt(3)
picture=doc.add_paragraph()
picture.paragraph_format.space_after=Pt(2)
picture.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=picture.add_run()
inline=run.add_picture(str(PNG),width=Inches(9.50))
inline._inline.docPr.set('descr','Retrospective replay architecture: saved Decision Tree and Isolation Forest exports, three-state OR merge, positive alerts opening declared case scopes, separate triggers and wider context, retrospective CTI, BM25 ATTACK reference selection, local Ollama inference and evidence review. The full saved VM3 Decision Tree test export is included; model test populations differ.')
doc.add_paragraph('Figure 1. Implemented replay of saved predictions into alert-linked case inputs. Wider context comes from processed flow records matching each declared scope. The full VM3 Decision Tree test export is included. Model test populations differ, and independent review of generated interpretation remains pending.','Caption')

sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
sec.orientation=WD_ORIENT.PORTRAIT
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.65);sec.right_margin=Inches(.65)

def para(s,style=None):
    return doc.add_paragraph(s,style)

def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,(c,h,w) in enumerate(zip(t.rows[0].cells,headers,widths)):
        c.width=Inches(w);c.text=h
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells=t.add_row().cells
        for c,value,w in zip(cells,values,widths):c.width=Inches(w);c.text=value
    borders=OxmlElement('w:tblBorders')
    for tag in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+tag);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    t._tbl.tblPr.append(borders)
    for idx,row in enumerate(t.rows):
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side,val in [('top','105'),('bottom','105'),('left','110'),('right','110')]:
                m=OxmlElement('w:'+side);m.set(qn('w:w'),val);m.set(qn('w:type'),'dxa');margins.append(m)
            pr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'243746' if idx==0 else ('F1F4F6' if idx%2==0 else 'FFFFFF'));pr.append(shade)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=1.04
                for r in p.runs:
                    r.font.size=Pt(9.5)
                    r.font.color.rgb=RGBColor(255,255,255) if idx==0 else RGBColor(0,0,0)
                    r.bold=(idx==0)
    return t

para('Stage inputs and outputs','Heading 1')
para('The 1 October replay connects saved detector decisions to case evidence and a local LLM request. Detection produces alerts; the later stages explain the available evidence and propose actions for review. The replay does not refit or rerun the machine-learning models.')
rows=[
('Data preparation','Existing processed VM2 / VM3 flow features and separate metadata and annotations.','Behavioural features and analysis identifiers. Recorded baseline features include port numbers. Endpoint IP addresses remain analysis metadata; dataset labels are separate training / evaluation targets.'),
('Decision Tree','Original labelled training data and recorded 80/20 evaluation protocol.','Full saved test exports for VM2 (16,091 rows) and VM3 (27,953 rows). The supplied VM3 export agrees with all 11 recorded examples. Other flow decisions remain unavailable.'),
('Isolation Forest','Recorded benign training / validation / test split and malicious evaluation flows.','Saved anomaly scores and binary decisions. The benign split is 70/15/15; its test population differs from the DT test population.'),
('OR combination','Predictions joined by dataset and flow row ID; source identities and protocols checked.','Positive if either available prediction is positive; negative only if both are negative; otherwise unknown. Model contributions and missing values are retained.'),
('Case preparation and CTI','Positive alerts, declared endpoint/date scopes, processed flows and saved CTI findings.','Exact trigger rows plus wider matching context. Counts and source hashes connect these records to the prompt. CTI is retrospective attribution context with recorded limitations.'),
('Reference selection','Observed service facet and pinned official ATT&CK Enterprise 19.2 definitions.','BM25-ranked candidate definitions. Retrieval can return irrelevant material and does not establish that any technique occurred.'),
('Local LLM','Alert summary, network aggregates, CTI, definitions, instructions and JSON fields.','Recorded Llama 3.1 8B responses via local Ollama. Remcos is regenerated after the full DT import; RedLine input and response are unchanged. Exact requests, raw responses and hashes are retained.'),
('Evidence review','Saved response, original flow evidence, CTI limitations and official definitions.','Documented supported, unsupported and uncertain claims. Valid JSON and a cited definition do not establish factual accuracy; independent team review remains pending.'),
]
table(['Stage','Input','Output'],rows,[1.25,2.30,3.65])

doc.add_page_break()
para('How the components fit together','Heading 1')
para('The Decision Tree classifies traffic from behavioural features; Isolation Forest flags observations outside its learned normal pattern. The replay joins their saved decisions on the same dataset and flow ID. Any positive opens an alert, including when the other model has no saved prediction. Two negatives are required for a negative result. Other missing combinations remain unknown.')
para('VM2 exports overlap on 1,941 flows: 1,940 DT-only positives and one IF-only positive. VM3 exports overlap on 15,794 flows: 15,782 DT-only positives, six positives from both models, three IF-only positives and three negatives from both. Both intersections contain only malicious-labelled records; benign test populations do not overlap. These counts describe model disagreement and cannot establish overall ensemble accuracy or false-positive rates.')
para('RedLine retains 590 DT-positive / IF-negative triggers and 2,755 context flows. Remcos now has 6,765 triggers and 33,744 context flows: 6,759 DT-positive / IF-negative, two DT-negative / IF-positive and four IF-positive with DT unavailable. Context is evidence for interpretation, not a detected-attack count. The three web-endpoint flows in the older 33,747-flow Remcos case remain outside this replay’s alerted endpoint scope.')
para('Both scopes were declared for retrospective studies. Earlier case experiments used preselected traffic and compared fixed references with BM25 retrieval; those responses retain that provenance. The new replay uses BM25 references and separate generated responses. Malware-family annotations do not select the replay’s triggers or context and do not enter its LLM prompts. Unknown-labelled records omitted by original preprocessing remain outside the available population.')
para('CTI adds dated endpoint relationships; ATT&CK supplies behaviour definitions. Neither proves that a named technique occurred in these flows. For example, TCP traffic, a service label and IP-level CTI do not establish HTTP command-and-control. The local LLM’s explanations and suggested actions are assessed against the supplied evidence before they can support a report claim.')

para('Implementation and reporting status','Heading 2')
status=[
('Merge and handoff',f'Implemented and checked: 21 boundary tests and {json.loads((BASE.parent / "outputs/verification.json").read_text())["source_and_output_checks_passed"]} source/output and generation-provenance reconciliations. Both cases retain exact trigger IDs, context rows and prompt hashes.'),
('Source provenance','Duc supplied all 27,953 original VM3 test predictions. IDs, labels and 11 recorded examples agree. The saved baseline notebook includes ports; the exact detector version and training features still need confirmation.'),
('Local interpretation','The corrected Remcos response states its endpoint scope but omits alert and traffic totals; T1095 remains unsupported. RedLine proposes five unsupported techniques. No mapping is accepted as confirmed in the accompanying assessments.'),
('Reporting limits','Independent team validation remains pending. The project demonstrates retrospective evidence interpretation, not unknown-attack discovery, live deployment or a validated automatic response system.'),
]
table(['Component','Evidence status'],status,[1.65,5.55])

para('Source records','Heading 2')
p=para('Project meeting transcript, 1 October 2026; saved model exports and source manifests from the VM2 / VM3 GitHub branches; original RemcosRAT and RedLine case experiments; PRT840 OR Handoff package, 1 October 2026, including verification.json, handoff traces, generation manifests and evidence review. Source commits and checksums are retained in the package.')
p.paragraph_format.space_after=Pt(3)
doc.core_properties.title='PRT840 System Architecture'
doc.core_properties.subject='Network detection and LLM assisted evidence interpretation'
doc.core_properties.author=''
doc.core_properties.keywords='PRT840, system architecture, network evidence, MITRE ATTACK, CTI, Ollama'
doc.save(DOCX)
print(DOCX)
print(PNG)
print(SVG)
